#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将 config/*.xlsx 导出为 config/*.json，供 Godot 运行时优先加载。
Excel 格式与现有 General/MapCell/Map 表结构一致，无需修改表头。
用法（在项目根目录）: python tools/export_config.py
"""
from __future__ import annotations

import json
import os
import sys


def _project_root() -> str:
    root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    return root


def _config_dir() -> str:
    return os.path.join(_project_root(), "config")


def _norm_portrait(path: str) -> str:
    if not path or not path.strip():
        return path
    path = path.strip().replace("\\", "/")
    if path.startswith("res://"):
        return path
    return "res://" + path.lstrip("/")


def _norm_map_cell(path: str) -> str:
    if not path or not path.strip():
        return path
    path = path.strip().replace("\\", "/").lstrip("/")
    if not path:
        return "res://Res/Map/"
    if path.startswith("res://"):
        return path
    if path.lower().startswith("res/map/"):
        return "res://" + path
    return "res://Res/Map/" + path


def export_general(config_dir: str) -> None:
    try:
        import openpyxl
    except ImportError:
        print("请先安装 openpyxl: pip install -r tools/requirements.txt", file=sys.stderr)
        sys.exit(1)

    path = os.path.join(config_dir, "General.xlsx")
    if not os.path.isfile(path):
        print(f"跳过（不存在）: {path}", file=sys.stderr)
        return

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    if not ws:
        wb.close()
        return

    out = {}
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or len(row) < 5:
            continue
        try:
            id_val = int(row[0]) if row[0] is not None else 0
            attack = int(row[1]) if row[1] is not None else 0
            attack_range = int(row[2]) if row[2] is not None else 0
            move_range = int(row[3]) if row[3] is not None else 0
            portrait = (row[4] or "").strip() if len(row) > 4 else ""
            ghp = int(row[5]) if len(row) > 5 and row[5] is not None else 1
            if ghp < 1:
                ghp = 1
        except (TypeError, ValueError):
            continue
        out[str(id_val)] = {
            "id": id_val,
            "attack": attack,
            "attackRange": attack_range,
            "moveRange": move_range,
            "portraitPath": _norm_portrait(portrait),
            "ghp": ghp,
        }
    wb.close()

    out_path = os.path.join(config_dir, "general.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    print(f"已导出: {out_path} ({len(out)} 条)")


def export_map_cell(config_dir: str) -> None:
    try:
        import openpyxl
    except ImportError:
        return

    path = os.path.join(config_dir, "MapCell.xlsx")
    if not os.path.isfile(path):
        print(f"跳过（不存在）: {path}", file=sys.stderr)
        return

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    if not ws:
        wb.close()
        return

    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 2:
            continue
        try:
            cell_id = int(row[0]) if row[0] is not None else 0
            res = (row[1] or "").strip() if len(row) > 1 else ""
            if not res:
                continue
            out[str(cell_id)] = _norm_map_cell(res)
        except (TypeError, ValueError):
            continue
    wb.close()

    out_path = os.path.join(config_dir, "map_cell.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    print(f"已导出: {out_path} ({len(out)} 条)")


def export_map_grid(config_dir: str) -> None:
    try:
        import openpyxl
    except ImportError:
        return

    path = os.path.join(config_dir, "Map.xlsx")
    if not os.path.isfile(path):
        print(f"跳过（不存在）: {path}", file=sys.stderr)
        return

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    if not ws:
        wb.close()
        return

    grid = []
    for r, row in enumerate(ws.iter_rows(min_row=2, max_row=6, values_only=True)):
        if r >= 5:
            break
        if not row or len(row) < 11:
            grid.append([0] * 10)
            continue
        line = []
        for c in range(1, 11):
            val = row[c]
            try:
                line.append(int(val) if val is not None else 0)
            except (TypeError, ValueError):
                line.append(0)
        grid.append(line)
    while len(grid) < 5:
        grid.append([0] * 10)
    wb.close()

    out_path = os.path.join(config_dir, "map_grid.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(grid, f, indent=2)
    print(f"已导出: {out_path}")


def main() -> None:
    root = _project_root()
    config_dir = os.path.join(root, "config")
    if not os.path.isdir(config_dir):
        print(f"config 目录不存在: {config_dir}", file=sys.stderr)
        sys.exit(1)
    export_general(config_dir)
    export_map_cell(config_dir)
    export_map_grid(config_dir)
    print("配置导出完成。运行时将优先加载 config/*.json。")


if __name__ == "__main__":
    main()
